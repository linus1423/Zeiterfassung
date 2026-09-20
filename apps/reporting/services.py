"""Auswertung und Export.

Aus den Zeiteinträgen werden Zeilen gebaut (je Eintrag oder verdichtet),
daraus entsteht eine Excel- oder CSV-Datei mit frei wählbaren Spalten.
"""

from __future__ import annotations

import csv
import io
from collections import OrderedDict
from collections.abc import Iterable, Iterator
from datetime import date, time

from django.utils import timezone

from apps.groups.closing import ClosedPeriods
from apps.groups.periods import MONTH_NAMES, WEEKDAY_NAMES, period_for
from apps.groups.permissions import readable_groups
from apps.tracking.models import TimeEntry
from apps.tracking.utils import day_bounds

from .columns import DATE, HHMM, HOURS, NUMBER, TEXT, TIME, Column

# Weiterhin unter den alten Namen erreichbar.
WEEKDAYS = WEEKDAY_NAMES
MONTHS = MONTH_NAMES

# Mit einem dieser Zeichen beginnt für Excel eine Formel. Die Werte kommen
# zum Teil aus freien Textfeldern (Notiz, Gruppen- und Tätigkeitsname), also
# werden sie in beiden Formaten entschärft.
_RISKY_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def query_entries(
    user,
    *,
    start: date,
    end: date,
    group_ids: Iterable[int] | None = None,
    user_ids: Iterable[int] | None = None,
    activity_ids: Iterable[int] | None = None,
):
    """Abgeschlossene Einträge im Zeitraum, begrenzt auf das, was der Nutzer lesen darf."""
    period_start, _ = day_bounds(start)
    _, period_end = day_bounds(end)

    queryset = (
        TimeEntry.objects.filter(end__isnull=False, start__gte=period_start, start__lt=period_end)
        .filter(group__in=readable_groups(user))
        .select_related("user", "group", "activity")
        .prefetch_related("breaks")
        .order_by("user__last_name", "user__first_name", "start")
    )
    if group_ids:
        queryset = queryset.filter(group_id__in=list(group_ids))
    if user_ids:
        queryset = queryset.filter(user_id__in=list(user_ids))
    if activity_ids:
        queryset = queryset.filter(activity_id__in=list(activity_ids))
    return queryset


def _base_row(entry: TimeEntry, closed: ClosedPeriods | None = None) -> dict:
    local_start = timezone.localtime(entry.start)
    local_end = timezone.localtime(entry.end) if entry.end else None
    iso_year, iso_week, _ = local_start.date().isocalendar()
    # Der Abrechnungszeitraum folgt dem Zyklus der Gruppe (Issue 8).
    period = period_for(local_start.date(), entry.group.month_start_day)
    is_closed = closed is not None and closed.is_closed(entry.group_id, local_start.date())
    return {
        "personnel_number": entry.user.personnel_number,
        "last_name": entry.user.last_name,
        "first_name": entry.user.first_name,
        "full_name": entry.user.full_name,
        "email": entry.user.email,
        "group": entry.group.name,
        "cost_center": entry.group.cost_center,
        "activity": entry.activity.name if entry.activity else "",
        "date": local_start.date(),
        "weekday": WEEKDAYS[local_start.weekday()],
        "week": iso_week,
        "month": period.month_label,
        "year": period.start.year,
        "period_label": period.label,
        "period_start": period.start,
        "period_end": period.end,
        "period_closed": "ja" if is_closed else "",
        "start": local_start.time().replace(second=0, microsecond=0),
        "end": local_end.time().replace(second=0, microsecond=0) if local_end else None,
        "break_seconds": entry.break_duration.total_seconds(),
        "work_seconds": entry.duration.total_seconds(),
        "entry_count": 1,
        "source": entry.get_source_display(),
        "incomplete": "ja" if entry.is_incomplete else "",
        "note": entry.note,
        "updated_at": timezone.localtime(entry.updated_at).strftime("%d.%m.%Y %H:%M"),
        "_iso_year": iso_year,
    }


def build_rows(
    entries: Iterable[TimeEntry], grouping: str, closed: ClosedPeriods | None = None
) -> list[dict]:
    """Baut die Exportzeilen in der gewünschten Verdichtung."""
    rows = [_base_row(entry, closed) for entry in entries]
    if grouping == "entry":
        return rows

    buckets: OrderedDict[tuple, dict] = OrderedDict()
    for row in rows:
        if grouping == "user_day":
            key = (row["email"], row["date"])
            keep = (
                "personnel_number",
                "last_name",
                "first_name",
                "full_name",
                "email",
                "group",
                "cost_center",
                "date",
                "weekday",
                "week",
                "month",
                "year",
                "period_label",
                "period_start",
                "period_end",
                "period_closed",
            )
        elif grouping == "user_month":
            # Der Schlüssel ist der Abrechnungszeitraum, nicht der Kalendermonat:
            # Gruppen mit eigenem Zyklus bleiben so getrennt (Issue 8).
            key = (row["email"], row["period_start"], row["period_end"])
            keep = (
                "personnel_number",
                "last_name",
                "first_name",
                "full_name",
                "email",
                "month",
                "year",
                "period_label",
                "period_start",
                "period_end",
                "period_closed",
            )
        elif grouping == "activity":
            key = (row["group"], row["activity"])
            keep = ("group", "cost_center", "activity")
        else:
            raise ValueError(f"Unbekannte Verdichtung: {grouping}")

        bucket = buckets.get(key)
        if bucket is None:
            bucket = {field: row.get(field, "") for field in keep}
            bucket.update({"break_seconds": 0.0, "work_seconds": 0.0, "entry_count": 0})
            buckets[key] = bucket
        bucket["break_seconds"] += row["break_seconds"]
        bucket["work_seconds"] += row["work_seconds"]
        bucket["entry_count"] += 1

    return list(buckets.values())


def _hours(seconds: float) -> float:
    return round(seconds / 3600, 2)


def _hhmm(seconds: float) -> str:
    total_minutes = int(round(seconds / 60))
    return f"{total_minutes // 60}:{total_minutes % 60:02d}"


def cell_value(row: dict, column: Column):
    """Rohwert einer Zelle, noch ohne Formatierung für ein bestimmtes Format."""
    if column.key == "hours":
        return _hours(row.get("work_seconds", 0.0))
    if column.key == "hhmm":
        return _hhmm(row.get("work_seconds", 0.0))
    if column.key == "break_hhmm":
        return _hhmm(row.get("break_seconds", 0.0))
    return row.get(column.key, "")


def total_row(rows: list[dict], columns: list[Column]) -> list:
    """Summenzeile: nur Dauer- und Anzahlspalten werden summiert."""
    work = sum(row.get("work_seconds", 0.0) for row in rows)
    pause = sum(row.get("break_seconds", 0.0) for row in rows)
    count = sum(row.get("entry_count", 0) for row in rows)

    values = []
    for index, column in enumerate(columns):
        if column.key == "hours":
            values.append(_hours(work))
        elif column.key == "hhmm":
            values.append(_hhmm(work))
        elif column.key == "break_hhmm":
            values.append(_hhmm(pause))
        elif column.key == "entry_count":
            values.append(count)
        elif index == 0:
            values.append("Summe")
        else:
            values.append("")
    return values


# --- Excel ------------------------------------------------------------------


def _force_text_cells(cells) -> None:
    """Hält Werte aus Textfeldern Text, auch wenn sie wie eine Formel aussehen.

    openpyxl macht aus einem Text mit führendem Gleichheitszeichen eine
    Formel. Die Notiz eines Zeiteintrags ist frei wählbar, also würde sie
    sonst in der Datei der Buchhaltung ausgeführt.
    """
    for cell in cells:
        if isinstance(cell.value, str) and cell.value[:1] in _RISKY_PREFIXES:
            cell.data_type = "s"


def to_xlsx(rows: list[dict], columns: list[Column], *, title: str = "Zeiten") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31] or "Zeiten"

    sheet.append([column.label for column in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"

    for row in rows:
        sheet.append([cell_value(row, column) for column in columns])
        _force_text_cells(sheet[sheet.max_row])

    if rows:
        summary = total_row(rows, columns)
        sheet.append(summary)
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
        _force_text_cells(sheet[sheet.max_row])

    for index, column in enumerate(columns, start=1):
        letter = get_column_letter(index)
        width = max(len(column.label) + 2, 12)
        sheet.column_dimensions[letter].width = min(width, 40)
        number_format = None
        if column.kind == DATE:
            number_format = "DD.MM.YYYY"
        elif column.kind == TIME:
            number_format = "HH:MM"
        elif column.kind == HOURS:
            number_format = "0.00"
        elif column.kind == NUMBER:
            number_format = "0"
        for cell in sheet[letter][1:]:
            if number_format:
                cell.number_format = number_format
            if column.kind in (HHMM, TEXT):
                cell.alignment = Alignment(horizontal="left")

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


# --- CSV --------------------------------------------------------------------


def csv_safe(value: str) -> str:
    """Schützt vor Formel-Einschleusung: Excel wertet führende = + - @ als Formel aus."""
    if value and value[0] in _RISKY_PREFIXES:
        return "'" + value
    return value


def format_scalar(value, column: Column, decimal_separator: str = ",") -> str:
    """Formatiert einen einzelnen Wert deutsch, egal aus welcher Zeile er stammt."""
    if value is None or value == "":
        return ""
    if column.kind == DATE and isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if column.kind == TIME and isinstance(value, time):
        return value.strftime("%H:%M")
    if column.kind == HOURS and isinstance(value, (int, float)):
        return f"{value:.2f}".replace(".", decimal_separator)
    return csv_safe(str(value))


def format_value(row: dict, column: Column, decimal_separator: str = ",") -> str:
    """Menschenlesbarer Text einer Zelle, für CSV und für die Vorschau."""
    return format_scalar(cell_value(row, column), column, decimal_separator)


def to_csv_rows(
    rows: list[dict],
    columns: list[Column],
    *,
    decimal_separator: str = ",",
    with_total: bool = True,
) -> Iterator[list[str]]:
    """Liefert die CSV-Zeilen einzeln, damit auch große Exporte gestreamt werden können."""
    yield [column.label for column in columns]
    for row in rows:
        yield [format_value(row, column, decimal_separator) for column in columns]
    if rows and with_total:
        yield [
            format_scalar(value, column, decimal_separator)
            for value, column in zip(total_row(rows, columns), columns, strict=True)
        ]


def csv_stream(
    rows: list[dict],
    columns: list[Column],
    *,
    delimiter: str = ";",
    decimal_separator: str = ",",
) -> Iterator[str]:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=delimiter, lineterminator="\r\n")
    for values in to_csv_rows(rows, columns, decimal_separator=decimal_separator):
        writer.writerow(values)
        buffer.seek(0)
        chunk = buffer.read()
        buffer.seek(0)
        buffer.truncate(0)
        yield chunk


def csv_bytes(
    rows: list[dict],
    columns: list[Column],
    *,
    delimiter: str = ";",
    decimal_separator: str = ",",
    with_bom: bool = True,
) -> Iterator[bytes]:
    """Wie csv_stream, aber als Bytes und mit genau einem BOM am Anfang."""
    first = True
    for chunk in csv_stream(
        rows, columns, delimiter=delimiter, decimal_separator=decimal_separator
    ):
        data = chunk.encode("utf-8")
        if first:
            first = False
            if with_bom:
                data = b"\xef\xbb\xbf" + data
        yield data
