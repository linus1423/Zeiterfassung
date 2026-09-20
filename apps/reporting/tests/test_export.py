import io
from datetime import datetime, time, timedelta

import pytest
from django.utils import timezone

from apps.groups.models import Activity, GroupMembership
from apps.reporting import services
from apps.reporting.columns import resolve
from apps.tracking.models import BreakEntry, TimeEntry


@pytest.fixture
def entries(member, group, activity):
    created = []
    # Feste Tageszeit statt "jetzt minus zwei Tage": je nach Uhrzeit des
    # Testlaufs liefe der Eintrag sonst über Mitternacht und würde für die
    # Auswertung auf zwei Tage aufgeteilt (Issue 32).
    first_day = timezone.localdate() - timedelta(days=2)
    for offset in (0, 1):
        start = timezone.make_aware(
            datetime.combine(first_day + timedelta(days=offset), time(8, 0))
        )
        entry = TimeEntry.objects.create(
            user=member, group=group, activity=activity, start=start, end=start + timedelta(hours=8)
        )
        BreakEntry.objects.create(
            time_entry=entry,
            start=start + timedelta(hours=4),
            end=start + timedelta(hours=4, minutes=30),
        )
        created.append(entry)
    return created


def test_accountant_sees_all_groups(entries, accountant, member):
    today = timezone.localdate()
    found = services.query_entries(accountant, start=today - timedelta(days=7), end=today)

    assert found.count() == len(entries)


def test_plain_member_sees_nothing_in_export(entries, member):
    today = timezone.localdate()

    assert services.query_entries(member, start=today - timedelta(days=7), end=today).count() == 0


def test_group_admin_sees_only_own_group(entries, group_admin, other_group, make_user):
    outsider = make_user("extern@example.com")
    GroupMembership.objects.create(user=outsider, group=other_group)
    foreign_activity = Activity.objects.create(group=other_group, name="Fremd")
    start = timezone.now() - timedelta(hours=3)
    TimeEntry.objects.create(
        user=outsider,
        group=other_group,
        activity=foreign_activity,
        start=start,
        end=start + timedelta(hours=2),
    )

    today = timezone.localdate()
    found = services.query_entries(group_admin, start=today - timedelta(days=7), end=today)

    assert {entry.group_id for entry in found} == {entries[0].group_id}


def test_build_rows_per_entry_and_per_day(entries, accountant):
    today = timezone.localdate()
    queryset = services.query_entries(accountant, start=today - timedelta(days=7), end=today)

    per_entry = services.build_rows(queryset, "entry")
    per_month = services.build_rows(queryset, "user_month")

    assert len(per_entry) == 2
    assert len(per_month) == 1
    assert per_month[0]["entry_count"] == 2
    assert per_month[0]["work_seconds"] == pytest.approx(2 * 7.5 * 3600)


def test_hours_use_german_decimal_separator(entries, accountant):
    today = timezone.localdate()
    rows = services.build_rows(
        services.query_entries(accountant, start=today - timedelta(days=7), end=today), "entry"
    )
    columns = resolve(["hours", "hhmm", "break_hhmm"])

    values = [services.format_value(rows[0], column) for column in columns]

    assert values == ["7,50", "7:30", "0:30"]


def test_csv_escapes_formula_injection():
    assert services.csv_safe("=1+1") == "'=1+1"
    assert services.csv_safe("@SUM(A1)") == "'@SUM(A1)"
    assert services.csv_safe("Montage") == "Montage"


def test_csv_stream_has_header_and_total(entries, accountant):
    today = timezone.localdate()
    rows = services.build_rows(
        services.query_entries(accountant, start=today - timedelta(days=7), end=today), "entry"
    )
    columns = resolve(["full_name", "hours"])

    text = "".join(services.csv_stream(rows, columns))
    lines = [line for line in text.splitlines() if line]

    assert lines[0] == "Name;Arbeitszeit (h)"
    assert lines[-1] == "Summe;15,00"


def test_xlsx_is_a_readable_workbook(entries, accountant):
    from openpyxl import load_workbook

    today = timezone.localdate()
    rows = services.build_rows(
        services.query_entries(accountant, start=today - timedelta(days=7), end=today), "entry"
    )
    columns = resolve(["full_name", "date", "hours"])

    content = services.to_xlsx(rows, columns)
    sheet = load_workbook(io.BytesIO(content)).active

    assert [cell.value for cell in sheet[1]] == ["Name", "Datum", "Arbeitszeit (h)"]
    assert sheet.max_row == len(rows) + 2  # Kopfzeile und Summenzeile
    assert sheet.freeze_panes == "A2"


def test_unknown_columns_fall_back_to_default():
    columns = resolve(["gibt-es-nicht"])

    assert [column.key for column in columns][:2] == ["personnel_number", "full_name"]


def test_csv_bytes_carries_the_bom_exactly_once(entries, accountant):
    today = timezone.localdate()
    rows = services.build_rows(
        services.query_entries(accountant, start=today - timedelta(days=7), end=today), "entry"
    )
    columns = resolve(["full_name", "hours"])

    content = b"".join(services.csv_bytes(rows, columns, with_bom=True))

    assert content.startswith(b"\xef\xbb\xbf")
    assert content.count(b"\xef\xbb\xbf") == 1
    assert not b"".join(services.csv_bytes(rows, columns, with_bom=False)).startswith(
        b"\xef\xbb\xbf"
    )
